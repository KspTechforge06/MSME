"""Document readers: turn uploaded bytes into plain text (+ optional tables).

Supported: PDF, DOCX, TXT/MD, CSV, XLSX.
OCR for scanned PDFs is intentionally out of scope (see project limitations).
"""
from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx", ".xls"}
TABULAR_EXTENSIONS = {".csv", ".xlsx", ".xls"}


class UnsupportedFileType(ValueError):
    """Raised when an upload has an extension we do not handle."""


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        parts.append(f"[page {index}]\n{text.strip()}")
    return "\n".join(parts).strip()


def _read_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_csv_text(data: bytes) -> str:
    text = _decode_text(data)
    rows = list(csv.reader(io.StringIO(text)))
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows).strip()


def _read_xlsx_text(data: bytes) -> str:
    import pandas as pd

    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
    parts: list[str] = []
    for name, frame in sheets.items():
        parts.append(f"[sheet: {name}]")
        parts.append(frame.to_string(index=False))
    return "\n".join(parts).strip()


def extract_text(filename: str, data: bytes) -> str:
    """Return plain text for a supported document, or raise UnsupportedFileType."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(data)
    if suffix == ".docx":
        return _read_docx(data)
    if suffix in {".txt", ".md"}:
        return _decode_text(data).strip()
    if suffix == ".csv":
        return _read_csv_text(data)
    if suffix in {".xlsx", ".xls"}:
        return _read_xlsx_text(data)
    raise UnsupportedFileType(f"Unsupported file type: {suffix or filename}")


def extract_tables(filename: str, data: bytes) -> list[dict[str, Any]]:
    """Return table rows as dicts for CSV/XLSX files, else an empty list.

    Deterministic validation uses these rows directly instead of relying on the
    language model to read multi-row ledgers.
    """
    suffix = Path(filename).suffix.lower()
    if suffix not in TABULAR_EXTENSIONS:
        return []

    import pandas as pd

    if suffix == ".csv":
        frame = pd.read_csv(io.BytesIO(data))
        frames = {None: frame}
    else:
        frames = pd.read_excel(io.BytesIO(data), sheet_name=None)

    rows: list[dict[str, Any]] = []
    for frame in frames.values():
        for record in frame.to_dict(orient="records"):
            cleaned: dict[str, Any] = {}
            for key, value in record.items():
                if value is None:
                    cleaned[str(key)] = None
                elif isinstance(value, float) and value != value:  # NaN
                    cleaned[str(key)] = None
                else:
                    cleaned[str(key)] = value
            rows.append(cleaned)
    return rows
